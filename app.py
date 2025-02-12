import os
from flask import Flask, request, jsonify
from pptx import Presentation
from transformers import pipeline

# Initialize Flask app
app = Flask(__name__)

# Path to save uploaded files
UPLOAD_FOLDER = 'uploads'
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# Load the summarization model
summarizer = pipeline("summarization", model="facebook/bart-large-cnn")

# Ensure the upload folder exists
if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

def extract_text_from_ppt(file_path):
    """
    Extracts text from the uploaded PPT file.
    """
    prs = Presentation(file_path)
    text = ''
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + '\n'
    return text

@app.route('/summarize', methods=['POST'])
def summarize_ppt():
    if 'file' not in request.files:
        return jsonify({"error": "No file part"}), 400
    
    file = request.files['file']
    
    if file.filename == '':
        return jsonify({"error": "No selected file"}), 400
    
    if file and file.filename.endswith('.pptx'):
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
        file.save(file_path)

        # Extract text from the PPT
        ppt_text = extract_text_from_ppt(file_path)

        # Summarize the text using the transformer model
        summary = summarizer(ppt_text, max_length=150, min_length=50, do_sample=False)

        return jsonify({"summary": summary[0]['summary_text']})

    else:
        return jsonify({"error": "Invalid file format. Please upload a .pptx file"}), 400

if __name__ == '__main__':
    app.run(debug=True)
